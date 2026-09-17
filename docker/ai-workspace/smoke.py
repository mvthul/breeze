"""Offline runtime contract: real document round trips and PDF rendering, not just imports."""
import importlib.metadata
import json
import tempfile
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from pypdf import PdfReader
import pdfplumber

with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    frame = pd.DataFrame({"value": [10, 20, 30]})
    assert int(np.sum(frame["value"].to_numpy())) == 60
    frame.to_excel(root / "analysis.xlsx", index=False, engine="openpyxl")
    workbook = load_workbook(root / "analysis.xlsx", read_only=True)
    assert list(workbook.active.values) == [("value",), (10,), (20,), (30,)]
    workbook.close()
    assert int(pd.read_excel(root / "analysis.xlsx")["value"].sum()) == 60
    document = Document()
    document.add_paragraph("Analysis total: 60")
    document.save(root / "analysis.docx")
    assert Document(root / "analysis.docx").paragraphs[0].text == "Analysis total: 60"
    pdf = root / "analysis.pdf"
    page = canvas.Canvas(str(pdf))
    page.drawString(72, 760, "Analysis total: 60")
    page.save()
    assert len(PdfReader(pdf).pages) == 1
    assert "Analysis total: 60" in PdfReader(pdf).pages[0].extract_text()
    with pdfplumber.open(pdf) as document:
        assert "Analysis total: 60" in document.pages[0].extract_text()
    subprocess.run(["pdftoppm", "-singlefile", "-scale-to", "600", "-png",
                    str(pdf), str(root / "analysis")], check=True, capture_output=True)
    assert (root / "analysis.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
print(json.dumps({"status": "passed", "total": 60, "versions": {
    name: importlib.metadata.version(name)
    for name in ["numpy", "pandas", "openpyxl", "python-docx", "reportlab", "pypdf", "pdfplumber"]
}}))
